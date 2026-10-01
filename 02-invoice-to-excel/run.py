"""
Processes every PDF invoice in a folder and creates a formatted Excel report.

Run:  python run.py                          (uses invoices/ -> output/invoice_register.xlsx)
      python run.py "D:\\Client Invoices" report.xlsx
"""
import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extractor import process_pdf

HERE = Path(__file__).parent
COLUMNS = {
    "file": "File", "vendor": "Vendor", "vendor_gstin": "Vendor GSTIN", "invoice_no": "Invoice No",
    "invoice_date": "Invoice Date", "taxable_value": "Taxable Value", "cgst": "CGST", "sgst": "SGST",
    "igst": "IGST", "total": "Invoice Total", "status": "Status", "issues": "Issues",
}
MONEY_COLS = ["Taxable Value", "CGST", "SGST", "IGST", "Invoice Total"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")


def style_sheet(ws, df: pd.DataFrame):
    """Make the sheet look professional: blue header, widths, ₹ format, freeze top row."""
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) if len(df) else len(str(col))
        ws.column_dimensions[letter].width = min(width + 3, 60)
        if col in MONEY_COLS:
            for cell in ws[letter][1:]:
                cell.number_format = "₹#,##0.00"
        if col == "Invoice Date":
            for cell in ws[letter][1:]:
                cell.number_format = "DD-MM-YYYY"

    if "Status" in df.columns:
        status_idx = list(df.columns).index("Status")
        for row in ws.iter_rows(min_row=2):
            if row[status_idx].value == "Needs review":
                for cell in row:
                    cell.fill = RED_FILL


def build_report(input_dir: Path, output_file: Path) -> pd.DataFrame:
    pdfs = sorted(input_dir.glob("*.pdf"))
    if not pdfs:
        raise SystemExit(f"No PDF files found in {input_dir}")

    rows = []
    for pdf in pdfs:
        try:
            rows.append(process_pdf(pdf))
        except Exception as e:  # a broken PDF shouldn't stop the whole batch
            rows.append({"file": pdf.name, "status": "Needs review", "issues": f"Could not read PDF: {e}"})

    df = pd.DataFrame(rows).reindex(columns=list(COLUMNS)).rename(columns=COLUMNS)
    df["Invoice Date"] = pd.to_datetime(df["Invoice Date"])

    ok = df[df["Status"] == "OK"]
    summary = (
        ok.groupby(["Vendor", "Vendor GSTIN"])[MONEY_COLS].sum()
        .assign(Invoices=ok.groupby(["Vendor", "Vendor GSTIN"]).size())
        .reset_index()
    )
    monthly = (
        ok.assign(Month=ok["Invoice Date"].dt.strftime("%Y-%m"))
        .groupby("Month")[MONEY_COLS].sum().reset_index()
    )
    monthly["Total ITC (Input Tax Credit)"] = monthly["CGST"] + monthly["SGST"] + monthly["IGST"]

    output_file.parent.mkdir(parents=True, exist_ok=True)
    sheets = {
        "All Invoices": df,
        "Needs Review": df[df["Status"] != "OK"],
        "Vendor Summary": summary,
        "Monthly GST": monthly,
    }
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for name, sheet_df in sheets.items():
            sheet_df.to_excel(writer, sheet_name=name, index=False)
            style_sheet(writer.sheets[name], sheet_df)
    return df


def main():
    parser = argparse.ArgumentParser(description="Convert GST invoice PDFs into an Excel register")
    parser.add_argument("input_dir", nargs="?", default=HERE / "invoices", type=Path)
    parser.add_argument("output_file", nargs="?", default=HERE / "output" / "invoice_register.xlsx", type=Path)
    args = parser.parse_args()

    df = build_report(args.input_dir, args.output_file)
    review = df[df["Status"] != "OK"]
    print(f"Processed {len(df)} invoices -> {args.output_file}")
    print(f"  OK: {len(df) - len(review)}   Needs review: {len(review)}")
    for _, r in review.iterrows():
        print(f"  ! {r['File']}: {r['Issues']}")


if __name__ == "__main__":
    main()
